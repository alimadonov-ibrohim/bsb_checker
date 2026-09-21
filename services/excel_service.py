"""
Excel Service — generate result workbooks with openpyxl.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from pathlib import Path
from typing import List, Dict, Any
import os

from services.paths import uploads_dir as default_uploads_dir


class ExcelService:
    def __init__(self, output_dir: str | None = None):
        self.output_dir = Path(output_dir) if output_dir else default_uploads_dir()
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def generate_results_excel(
        self,
        test_info: Dict[str, Any],
        students_data: List[Dict[str, Any]],
        filename: str | None = None,
    ) -> str:
        """
        students_data item:
        {
          "name": "Ali Valiyev",
          "class_name": "9-A",
          "correct_count": 27,
          "incorrect_count": 3,
          "uncertain_count": 0,
          "percentage": 90.0,
          "score": 27,
          "total_questions": 30,
          "details": {"1": {"student": "A", "correct": "A", "status": "correct"}, ...}
        }
        """
        wb = Workbook()

        # ========== Sheet 1: Summary ==========
        ws1 = wb.active
        ws1.title = "Natijalar"

        headers1 = [
            "№", "O‘quvchi", "Sinf", "Jami savol",
            "To‘g‘ri", "Noto‘g‘ri", "Noaniq", "Foiz", "Ball"
        ]
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="2E7D32")
        thin = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        for col, h in enumerate(headers1, 1):
            cell = ws1.cell(1, col, h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin

        for idx, s in enumerate(students_data, 1):
            row = [
                idx,
                s.get("name", f"O‘quvchi #{idx}"),
                s.get("class_name") or test_info.get("class_name", ""),
                s.get("total_questions", test_info.get("question_count", 0)),
                s.get("correct_count", 0),
                s.get("incorrect_count", 0),
                s.get("uncertain_count", 0),
                s.get("percentage", 0),
                s.get("score", 0),
            ]
            for col, val in enumerate(row, 1):
                cell = ws1.cell(idx + 1, col, val)
                cell.border = thin
                cell.alignment = Alignment(horizontal="center")

        for col in range(1, len(headers1) + 1):
            ws1.column_dimensions[get_column_letter(col)].width = 14
        ws1.column_dimensions["B"].width = 22

        # ========== Sheet 2: Detailed answers ==========
        ws2 = wb.create_sheet("Batafsil javoblar")

        # Collect all question numbers
        all_q = set()
        for s in students_data:
            all_q.update(s.get("details", {}).keys())
        sorted_q = sorted(all_q, key=lambda x: int(x) if str(x).isdigit() else 0)

        headers2 = ["O‘quvchi"] + [f"{q}-savol" for q in sorted_q] + ["To‘g‘ri", "Noto‘g‘ri", "Foiz"]
        for col, h in enumerate(headers2, 1):
            cell = ws2.cell(1, col, h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = thin

        green_fill = PatternFill("solid", fgColor="C8E6C9")
        red_fill = PatternFill("solid", fgColor="FFCDD2")
        yellow_fill = PatternFill("solid", fgColor="FFF9C4")

        for row_idx, s in enumerate(students_data, 2):
            ws2.cell(row_idx, 1, s.get("name", "")).border = thin
            details = s.get("details", {})
            for col_idx, q in enumerate(sorted_q, 2):
                d = details.get(str(q), {})
                student_ans = d.get("student")
                status = d.get("status", "uncertain")
                display = student_ans if student_ans else "—"
                cell = ws2.cell(row_idx, col_idx, display)
                cell.border = thin
                cell.alignment = Alignment(horizontal="center")
                if status == "correct":
                    cell.fill = green_fill
                elif status == "incorrect":
                    cell.fill = red_fill
                else:
                    cell.fill = yellow_fill

            last_cols = [
                s.get("correct_count", 0),
                s.get("incorrect_count", 0),
                s.get("percentage", 0),
            ]
            base = 2 + len(sorted_q)
            for i, val in enumerate(last_cols):
                cell = ws2.cell(row_idx, base + i, val)
                cell.border = thin
                cell.alignment = Alignment(horizontal="center")

        ws2.column_dimensions["A"].width = 22
        for col in range(2, len(headers2) + 1):
            ws2.column_dimensions[get_column_letter(col)].width = 10

        # Meta sheet
        ws3 = wb.create_sheet("Test ma'lumoti")
        meta = [
            ("Fan", test_info.get("subject", "")),
            ("Sinf", test_info.get("class_name", "")),
            ("Test nomi", test_info.get("test_name", "")),
            ("Test turi", test_info.get("test_type", "")),
            ("Savollar soni", test_info.get("question_count", "")),
            ("O‘quvchilar soni", len(students_data)),
        ]
        for i, (k, v) in enumerate(meta, 1):
            ws3.cell(i, 1, k).font = Font(bold=True)
            ws3.cell(i, 2, v)
        ws3.column_dimensions["A"].width = 20
        ws3.column_dimensions["B"].width = 30

        if not filename:
            safe_name = f"natija_{test_info.get('id', 'test')}.xlsx"
        else:
            safe_name = filename

        out_path = self.output_dir / safe_name
        wb.save(str(out_path))
        return str(out_path)
